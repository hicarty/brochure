"""
Integrate Proposal Theme — reusable python-pptx helpers.

Import this into any new deck-generation script to keep colours, type and
layout consistent with the Integrate brand system. Pair it with
integrate_theme.md (the plain-language spec) when briefing a new proposal.

    from integrate_theme import *

    prs = new_presentation()
    slide = add_slide(prs, DARK)
    eyebrow(slide, "THE OPPORTUNITY")
    title(slide, "The Gap We're Closing")
    footer_logo(slide, "assets/logo_white.png")
    prs.save("Proposal.pptx")
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- palette --
ORANGE    = RGBColor(0xF0, 0x55, 0x39)
DARK      = RGBColor(0x24, 0x23, 0x28)
CARD_DARK = RGBColor(0x1B, 0x1A, 0x1F)
BLUE      = RGBColor(0x28, 0x6F, 0xBF)
TEAL      = RGBColor(0x2A, 0x6C, 0x86)
GREEN     = RGBColor(0x8F, 0xBF, 0x5C)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GREY      = RGBColor(0xC9, 0xC7, 0xC5)
GREY_DARK = RGBColor(0x9A, 0x97, 0x93)
DIVIDER   = RGBColor(0x3A, 0x39, 0x40)

# ----------------------------------------------------------------- type ---
FONT_HEAD = "Arial Black"
FONT_BODY = "Arial"

# --------------------------------------------------------------- canvas ---
PAGE_W = Inches(13.333)   # 16:9 widescreen — matches the 960x540pt source
PAGE_H = Inches(7.5)


def new_presentation():
    """A blank 16:9 presentation sized to match the Integrate decks."""
    prs = Presentation()
    prs.slide_width = PAGE_W
    prs.slide_height = PAGE_H
    return prs


def add_slide(prs, bg_color=ORANGE):
    """Blank slide with a solid background fill."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg_color
    return slide


# ------------------------------------------------------------- textboxes --
def _textbox(slide, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return box, tf


def _style(run, font=FONT_BODY, size=12, color=WHITE, bold=False, italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic


def page_num(slide, n, color=WHITE):
    box, tf = _textbox(slide, 0.4, 0.28, 0.8, 0.35)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = f"{n:02d}"
    _style(r, size=12, color=color, bold=True)


def eyebrow(slide, text, x=0.55, y=0.35, w=9, color=WHITE):
    box, tf = _textbox(slide, x, y, w, 0.35)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text.upper()
    _style(r, size=12, color=color, bold=True)


def title(slide, text, x=0.55, y=0.65, w=10.8, h=0.95, size=32, color=WHITE):
    box, tf = _textbox(slide, x, y, w, h)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    _style(r, font=FONT_HEAD, size=size, color=color, bold=True)


def subtitle(slide, text, x=0.55, y=1.55, w=10.8, h=0.6, size=14, color=GREY):
    box, tf = _textbox(slide, x, y, w, h)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    _style(r, size=size, color=color)


def body_text(slide, text, x, y, w, h, size=12, color=GREY, bold=False,
              italic=False, align=PP_ALIGN.LEFT, line_spacing=1.25):
    box, tf = _textbox(slide, x, y, w, h)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    _style(r, size=size, color=color, bold=bold, italic=italic)
    return box


def footer_logo(slide, path, x=11.55, y=6.95, w=1.35, h=0.293):
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))


# ----------------------------------------------------------------- shapes --
def rounded_rect(slide, x, y, w, h, fill=CARD_DARK, radius=0.08):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    return shp


def circle(slide, x, y, d, fill=ORANGE):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    return shp


def divider_line(slide, x, y, w, color=DIVIDER, weight_pt=0.75):
    shp = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    shp.line.color.rgb = color
    shp.line.width = Pt(weight_pt)
    return shp


def mission_blob(slide, x=7.15, y=0.55, w=5.6, h=6.4, fill=BLUE, radius=0.35):
    """The large soft rounded panel used for mission/ask statements."""
    return rounded_rect(slide, x, y, w, h, fill=fill, radius=radius)


# ------------------------------------------------------------ composites --
def card(slide, x, y, w, h, glyph, card_title, card_body,
         icon_color=ORANGE, card_fill=CARD_DARK):
    """Icon-circle + title + body card — the standard service/audience tile."""
    rounded_rect(slide, x, y, w, h, fill=card_fill)
    d = 0.58
    circ = circle(slide, x + w / 2 - d / 2, y + 0.3, d, fill=icon_color)
    tf = circ.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = glyph
    _style(r, size=18, color=WHITE, bold=True)
    body_text(slide, card_title, x + 0.18, y + 1.05, w - 0.36, 0.6,
               size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.05)
    body_text(slide, card_body, x + 0.2, y + 1.7, w - 0.4, h - 1.9,
               size=10.5, color=GREY, align=PP_ALIGN.CENTER, line_spacing=1.18)


def stat_chip(slide, x, y, w, h, value, label, value_color=GREEN):
    rounded_rect(slide, x, y, w, h, fill=CARD_DARK, radius=0.06)
    body_text(slide, value, x + 0.1, y + 0.12, w - 0.2, h * 0.55,
               size=24, color=value_color, bold=True, align=PP_ALIGN.CENTER)
    body_text(slide, label, x + 0.1, y + h * 0.62, w - 0.2, h * 0.35,
               size=10.5, color=GREY, align=PP_ALIGN.CENTER)


# Glyph shorthand used across Integrate decks (plain-shape "icons" —
# no icon library dependency, matches the originals' look):
GLYPH_ARROW  = "\u2192"
GLYPH_SQUARE = "\u25A0"
GLYPH_DOT    = "\u25CF"
GLYPH_TRI    = "\u25B2"
GLYPH_RING   = "\u25CE"
GLYPH_STAR   = "\u2726"
